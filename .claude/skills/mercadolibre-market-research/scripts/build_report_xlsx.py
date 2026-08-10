#!/usr/bin/env python3
"""Construye el Excel final de un estudio: dos hojas, "Resumen por nicho" y
"Productos", con formato de tabla, moneda, encabezados congelados y enlaces.

Por qué existe: entregar un CSV a un usuario que lo va a abrir en el celular
es ilegible (se ve como texto plano lleno de comas). Un .xlsx con columnas
en español, formato de moneda y una hoja de resumen arriba de los datos crudos
se abre bien en cualquier app de hojas de cálculo (Excel, Google Sheets, la
app nativa del teléfono) y comunica el hallazgo, no solo los datos.

Uso:
    python3 build_report_xlsx.py resumen_nichos.json productos.json salida.xlsx \
        --moneda CLP --titulo "Productos de baño en MercadoLibre Chile" \
        --fecha "10 de agosto de 2026" --sitio "MLC"

Formato EXACTO esperado de resumen_nichos.json — una lista de dicts, uno por
nicho, con TODAS estas claves (no un subconjunto; si renombras una clave la
columna correspondiente sale vacía sin ningún error visible, así que respeta
los nombres tal cual):
    nicho, score, ofertas, vendedores, uds_publ, uds_tot,
    p50, pmin, pmax, peso, catalogo, full, desc, disp

Formato EXACTO esperado de productos.json — la salida de parse_cards.py
(title, link, seller, sold, sold_txt, rating, price, list_price, discount,
free_ship, full, top_seller), más los dos campos que agregas tú al
clasificar cada producto:
    nicho     (obligatorio: a qué nicho pertenece)
    peso_kg   (obligatorio: peso estimado usado en el modelo de envío)
    src_cat   (opcional: subcategoría de origen, si la tienes)

Este script valida que las columnas críticas no queden vacías por un typo de
nombre de clave — es el error más común al armar estos JSON a mano en una
sesión nueva — y si detecta una columna completamente vacía, además de avisar
por stderr, deja una fila de advertencia roja al principio de la hoja para
que salte a la vista aunque no se haya leído la consola.
"""
import argparse
import json
import re
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Arial"
NAVY = "1F3B4D"
LIGHT = "F4F5F3"
WHITE = "FFFFFF"
ALERT = "B3261E"
ALERT_BG = "FBE9E7"

header_font = Font(name=FONT, bold=True, color=WHITE, size=10)
header_fill = PatternFill("solid", fgColor=NAVY)
title_font = Font(name=FONT, bold=True, size=14, color=NAVY)
sub_font = Font(name=FONT, italic=True, size=9, color="65736F")
body_font = Font(name=FONT, size=10)
alert_font = Font(name=FONT, bold=True, size=10, color=ALERT)
thin = Side(style="thin", color="D9DEDA")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

RESUMEN_COLS = [
    ("Nicho", "nicho", 34, None),
    ("Score", "score", 9, "0.0"),
    ("Ofertas", "ofertas", 9, "#,##0"),
    ("Competidores", "vendedores", 12, "#,##0"),
    ("Uds/publicación", "uds_publ", 13, "#,##0"),
    ("Uds. totales", "uds_tot", 12, "#,##0"),
    ("Precio mediano", "p50", 13, "MONEDA"),
    ("Precio mín.", "pmin", 12, "MONEDA"),
    ("Precio máx.", "pmax", 12, "MONEDA"),
    ("Peso (kg)", "peso", 9, "0.0"),
    ("% en catálogo", "catalogo", 11, '0"%"'),
    ("% Full", "full", 9, '0"%"'),
    ("Desc. mediano", "desc", 11, '0"%"'),
    ("% disponible tras costos", "disp", 14, '0"%"'),
]
# claves que deberían tener valor en (casi) todas las filas; si TODAS las filas
# de una de estas quedan vacías, es casi seguro un bug de nombre de campo.
RESUMEN_CRITICAL = ["vendedores", "uds_publ", "uds_tot", "p50", "peso", "catalogo", "full", "disp"]

PRODUCTO_COLS = [
    ("Nicho", "nicho", 30, None, False),
    ("Subcategoría", "src_cat", 20, None, False),
    ("Producto", "title", 55, None, False),
    ("Precio", "price", 13, "MONEDA", False),
    ("Precio de lista", "list_price", 15, "MONEDA", False),
    ("Descuento", "_desc_pct", 11, "0%", False),
    ("Unidades vendidas", "sold", 13, '"+"#,##0', False),
    ("Calificación", "rating", 11, '0.0"★"', False),
    ("Vendedor", "seller", 20, None, False),
    ("Full", "_full", 8, None, False),
    ("Más vendido", "_top", 11, None, False),
    ("Peso est. (kg)", "peso_kg", 11, "0.0", False),
    ("Enlace", "link", 12, None, False),
]
# "seller" queda vacío en ~50% de las tarjetas de forma normal (MercadoLibre no
# siempre lo muestra) y "src_cat" es opcional — ninguno de los dos se valida.
PRODUCTO_CRITICAL = ["nicho", "price", "peso_kg"]


def style_header(ws, row, cols, moneda_labels=()):
    for j, col in enumerate(cols, start=1):
        label = col[0]
        if col[1] in moneda_labels:
            label = f"{label} ({moneda_labels[col[1]]})"
        c = ws.cell(row=row, column=j, value=label)
        c.font, c.fill, c.border = header_font, header_fill, border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def columnas_vacias(rows, keys, umbral=0.8):
    """Devuelve las claves donde >= `umbral` de las filas tienen valor None."""
    if not rows:
        return []
    vacias = []
    for k in keys:
        faltan = sum(1 for r in rows if r.get(k) is None)
        if faltan / len(rows) >= umbral:
            vacias.append(k)
    return vacias


def insertar_alerta(ws, row, ncols, columnas):
    """Escribe la alerta en una fila YA reservada (la fila 3, el espaciador
    entre el subtítulo y el encabezado) — nunca inserta filas nuevas, porque
    insert_rows() no reubica el rango de una Table ya creada y deja el
    archivo con la tabla apuntando a la fila equivocada."""
    label_map = {k: lab for lab, k, *_ in RESUMEN_COLS + [(c[0], c[1]) for c in PRODUCTO_COLS]}
    nombres = ", ".join(label_map.get(k, k) for k in columnas)
    c = ws.cell(row=row, column=1, value=f"⚠ Columnas vacías en casi todas las filas — revisa el JSON de origen: {nombres}")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c.font, c.alignment = alert_font, Alignment(vertical="center")
    for j in range(1, ncols + 1):
        ws.cell(row=row, column=j).fill = PatternFill("solid", fgColor=ALERT_BG)


def build(resumen_path, productos_path, out_path, moneda="CLP", titulo="Estudio de mercado", fecha=None, sitio=None):
    money_fmt = "$#,##0" if moneda in ("CLP", "COP") else f'$#,##0"{" " + moneda if moneda != "USD" else ""}"'

    wb = Workbook()

    # ---- Hoja 1: Resumen por nicho ----
    ws1 = wb.active
    ws1.title = "Resumen por nicho"
    resumen = json.load(open(resumen_path, encoding="utf-8"))
    resumen = sorted(resumen, key=lambda r: -(r.get("score") or 0))

    vacias1 = columnas_vacias(resumen, RESUMEN_CRITICAL)
    if vacias1:
        print(f"ADVERTENCIA: en resumen_nichos.json, estas claves están vacías en casi todas las filas: {vacias1}", file=sys.stderr)
        print("Revisa que el JSON use exactamente los nombres de campo documentados en el docstring de este script.", file=sys.stderr)

    ws1["A1"] = f"{titulo} — resumen por nicho"
    ws1["A1"].font = title_font
    partes = []
    if fecha:
        partes.append(f"Captura: {fecha}")
    if sitio:
        partes.append(f"Sitio {sitio}")
    partes.append(f"precios en {moneda}")
    partes.append("Score = demanda x (1/competencia) x disponibilidad de catálogo x margen")
    ws1["A2"] = " · ".join(partes)
    ws1["A2"].font = sub_font
    ws1.merge_cells(f"A1:{get_column_letter(len(RESUMEN_COLS))}1")
    ws1.merge_cells(f"A2:{get_column_letter(len(RESUMEN_COLS))}2")

    header_row = 4
    money_cols1 = {k: moneda for _l, k, _w, fmt in RESUMEN_COLS if fmt == "MONEDA"}
    style_header(ws1, header_row, RESUMEN_COLS, money_cols1)
    for i, r in enumerate(resumen, start=header_row + 1):
        for j, (label, key, width, fmt) in enumerate(RESUMEN_COLS, start=1):
            c = ws1.cell(row=i, column=j, value=r.get(key))
            c.font, c.border = body_font, border
            if fmt:
                c.number_format = money_fmt if fmt == "MONEDA" else fmt
            c.alignment = Alignment(
                vertical="center", wrap_text=(key == "nicho"), horizontal="left" if key == "nicho" else "center"
            )
        if i % 2 == 0:
            for j in range(1, len(RESUMEN_COLS) + 1):
                ws1.cell(row=i, column=j).fill = PatternFill("solid", fgColor=LIGHT)
    for j, (_l, _k, width, _f) in enumerate(RESUMEN_COLS, start=1):
        ws1.column_dimensions[get_column_letter(j)].width = width
    if resumen:
        ws1.add_table(
            Table(
                displayName="ResumenNichos",
                ref=f"A{header_row}:{get_column_letter(len(RESUMEN_COLS))}{header_row + len(resumen)}",
                tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=False),
            )
        )
    ws1.row_dimensions[header_row].height = 28
    if vacias1:
        insertar_alerta(ws1, 3, len(RESUMEN_COLS), vacias1)
    ws1.freeze_panes = f"A{header_row + 1}"

    # ---- Hoja 2: Productos ----
    ws2 = wb.create_sheet("Productos")
    productos = json.load(open(productos_path, encoding="utf-8"))
    for r in productos:
        m = re.search(r"\d+", r.get("discount") or "")
        r["_desc_pct"] = int(m.group(0)) / 100 if m else None
        r["_full"] = "Sí" if r.get("full") else "No"
        r["_top"] = "Sí" if r.get("top_seller") else "No"

    vacias2 = columnas_vacias(productos, PRODUCTO_CRITICAL)
    if vacias2:
        print(f"ADVERTENCIA: en productos.json, estas claves están vacías en casi todas las filas: {vacias2}", file=sys.stderr)

    productos.sort(key=lambda r: (r.get("nicho", ""), -(r.get("sold") or 0), r.get("price") or 0))

    ws2["A1"] = f"{titulo} — detalle de productos"
    ws2["A1"].font = title_font
    sub2 = f"{len(productos)} publicaciones únicas · superficie de ofertas por categoría"
    if fecha:
        sub2 += f" · capturado el {fecha}"
    ws2["A2"] = sub2
    ws2["A2"].font = sub_font

    header_row2 = 4
    money_cols2 = {k: moneda for _l, k, _w, fmt, *_r in PRODUCTO_COLS if fmt == "MONEDA"}
    style_header(ws2, header_row2, PRODUCTO_COLS, money_cols2)
    for i, r in enumerate(productos, start=header_row2 + 1):
        for j, (label, key, width, fmt, wrap) in enumerate(PRODUCTO_COLS, start=1):
            if key == "link":
                c = ws2.cell(row=i, column=j, value="Ver ficha ↗")
                if r.get("link"):
                    c.hyperlink = r["link"]
                c.font = Font(name=FONT, size=10, color="0D6B5F", underline="single")
            else:
                c = ws2.cell(row=i, column=j, value=r.get(key))
                c.font = body_font
                if fmt:
                    c.number_format = money_fmt if fmt == "MONEDA" else fmt
            c.border = border
            c.alignment = Alignment(
                vertical="center",
                wrap_text=wrap,
                horizontal="left" if key in ("nicho", "src_cat", "title", "seller") else "center",
            )
        if i % 2 == 0:
            for j in range(1, len(PRODUCTO_COLS) + 1):
                ws2.cell(row=i, column=j).fill = PatternFill("solid", fgColor=LIGHT)
    for j, (_l, _k, width, _f, _w) in enumerate(PRODUCTO_COLS, start=1):
        ws2.column_dimensions[get_column_letter(j)].width = width
    if productos:
        ws2.add_table(
            Table(
                displayName="Productos",
                ref=f"A{header_row2}:{get_column_letter(len(PRODUCTO_COLS))}{header_row2 + len(productos)}",
                tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=False),
            )
        )
    ws2.row_dimensions[header_row2].height = 30
    if vacias2:
        insertar_alerta(ws2, 3, len(PRODUCTO_COLS), vacias2)
    ws2.freeze_panes = f"D{header_row2 + 1}"

    wb.save(out_path)
    estado = "CON ADVERTENCIAS" if (vacias1 or vacias2) else "ok"
    print(f"guardado {out_path} · {len(resumen)} nichos · {len(productos)} productos · {estado}")
    return bool(vacias1 or vacias2)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("resumen_json")
    ap.add_argument("productos_json")
    ap.add_argument("salida_xlsx")
    ap.add_argument("--moneda", default="CLP")
    ap.add_argument("--titulo", default="Estudio de mercado en MercadoLibre")
    ap.add_argument("--fecha", default=None, help='ej. "10 de agosto de 2026"')
    ap.add_argument("--sitio", default=None, help='ej. "MLC"')
    args = ap.parse_args()
    hubo_advertencias = build(
        args.resumen_json, args.productos_json, args.salida_xlsx,
        args.moneda, args.titulo, args.fecha, args.sitio,
    )
    sys.exit(1 if hubo_advertencias else 0)
